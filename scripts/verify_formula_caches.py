"""Compare every cached formula result to independent recalculation, preserving XLSX features."""
import json
import math
from openpyxl import load_workbook
from src.common import ROOT

path=ROOT/'output/Lutenitsa_Factory_Feasibility_2015_2025.xlsx'
formulas=load_workbook(path,data_only=False)
values=load_workbook(path,data_only=True)
computed=json.loads((ROOT/'output/qa/recalculated_cells.json').read_text(encoding='utf-8'))
errors=[]
count=0
for record in computed:
    sheet=formulas[record['sheet']]
    cache=values[record['sheet']]
    # Every generated used range starts at B2; assert the title anchor explicitly.
    assert record['values'][0][0]==sheet['B2'].value
    for row in sheet:
        for cell in row:
            if cell.data_type!='f':
                continue
            count+=1
            expected=record['values'][cell.row-2][cell.column-2]
            saved=cache[cell.coordinate].value
            if isinstance(expected,(float,int)) and isinstance(saved,(float,int)):
                same=math.isclose(expected,saved,rel_tol=1e-10,abs_tol=.00001)
            else:
                same=expected==saved
            if not same:
                errors.append(dict(cell=sheet.title+'!'+cell.coordinate,saved=saved,recalculated=expected))
assert not errors,errors
result=dict(status='PASS',formula_count=count,errors=errors,cache_policy='Every formula cache matched independent recalculation. Original XlsxWriter file preserved, including native charts/tables/validations.')
(ROOT/'reports/formula_cache_verification.json').write_text(json.dumps(result,indent=2),encoding='utf-8')
print(f'All {count} cached formulas match independent recalculation.')
