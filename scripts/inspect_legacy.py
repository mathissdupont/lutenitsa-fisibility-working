"""Inspect legacy workbooks without writing to original files."""
from pathlib import Path
import hashlib
import json
from openpyxl import load_workbook

root = Path(__file__).resolve().parents[1]
(root / 'reports').mkdir(exist_ok=True)
inventory = []
for path in sorted(root.glob('*')):
    if path.is_file():
        inventory.append({'file': path.name, 'sha256': hashlib.sha256(path.read_bytes()).hexdigest(), 'bytes': path.stat().st_size})
details = {}
for path in root.glob('*.xlsx'):
    formulas = load_workbook(path, data_only=False)
    values = load_workbook(path, data_only=True)
    details[path.name] = {}
    for ws in formulas:
        rows = []
        for row in ws:
            populated = [{'cell': c.coordinate, 'formula_or_value': c.value, 'cached_value': values[ws.title][c.coordinate].value} for c in row if c.value is not None]
            if populated:
                rows.append(populated)
        details[path.name][ws.title] = rows
(root / 'reports/legacy_inventory.json').write_text(json.dumps(inventory, ensure_ascii=False, indent=2), encoding='utf-8')
(root / 'reports/legacy_workbooks.json').write_text(json.dumps(details, ensure_ascii=False, indent=2, default=str), encoding='utf-8')
for name, sheets in details.items():
    print(name, list(sheets))
    if name.endswith('(1).xlsx'):
        continue
    for sheet, rows in sheets.items():
        print(sheet)
        for row in rows[:18]:
            print([(c['cell'], c['formula_or_value'], c['cached_value']) for c in row])
