"""Print bounded official-source excerpts for extraction design."""
from pathlib import Path
from openpyxl import load_workbook
from pypdf import PdfReader
root=Path(__file__).resolve().parents[1]
for filename in ['ilcemesafe.xlsx','ilmesafe.xlsx']:
    workbook=load_workbook(root/'data/raw/official'/filename,read_only=True,data_only=True)
    print(filename,workbook.sheetnames)
    for sheet in workbook:
        print(sheet.title,sheet.max_row,sheet.max_column)
        for row in sheet.iter_rows(min_row=1,max_row=7,max_col=20,values_only=True):
            print(row)
    workbook.close()
for filename in ['tomato_report_2024.pdf','guide_6.pdf','guide_7.pdf','guide_8.pdf']:
    reader=PdfReader(root/'data/raw/official'/filename)
    print(filename)
    for i,page in enumerate(reader.pages):
        text=page.extract_text() or ''
        if (filename.startswith('guide') and 'Salçalık' in text) or (filename.startswith('tomato') and 'Ek Tablo 11.' in text):
            print('PAGE',i+1,text[:6500])
