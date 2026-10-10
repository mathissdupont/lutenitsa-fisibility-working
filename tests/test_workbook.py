"""Validate the actual deliverable without rewriting it."""
import json
import pytest
from openpyxl import load_workbook
from src.common import ROOT
from src.reporting import SHEETS

@pytest.fixture(scope='module')
def books():
    path=ROOT/'output/Lutenitsa_Factory_Feasibility_2015_2025.xlsx'
    return load_workbook(path,data_only=False),load_workbook(path,data_only=True)

def test_workbook_structure_and_controls(books):
    formulas,values=books
    assert formulas.sheetnames==SHEETS
    assert sum(len(s.tables) for s in formulas)>=15
    assert len(formulas['PRODUCT_TRENDS']._charts)==1
    assert formulas['HISTORICAL_PRODUCTION'].freeze_panes=='C6'
    assert len(formulas['ASSUMPTIONS'].data_validations.dataValidation)>=3
    assert len(formulas['FE401_DECISION_MATRIX'].data_validations.dataValidation)==2
    assert values['DASHBOARD']['C11'].value==2500000
    assert values['DASHBOARD']['C12'].value==5000000

def test_workbook_missing_finance_and_no_errors(books):
    formulas,values=books
    count=0
    for sheet in formulas:
        for row in sheet:
            for cell in row:
                assert values[sheet.title][cell.coordinate].data_type!='e'
                if cell.data_type=='f':
                    count+=1
                    assert values[sheet.title][cell.coordinate].value is not None
    assert count>800
    assert values['PROFITABILITY_NPV']['D6'].value=='Eksik'
    assert values['ASSUMPTIONS']['C13'].value is None
    assert values['OPERATING_COSTS']['D30'].value=='Eksik'

def test_no_demo_results_in_main_workbook(books):
    _,values=books
    for sheet in values:
        for row in sheet:
            for cell in row:
                assert cell.value not in [83674650,84611000,85774000]
    assert values['FORECASTS_2026_2030']['G6'].value is None

def test_regression_not_fabricated():
    from src.common import read_csv
    assert (read_csv('data/processed/model_validation.csv').status=='NOT_ESTIMABLE').all()
    assert read_csv('data/processed/regression_coefficients.csv').empty
    assert read_csv('data/processed/forecasts.csv').baseline.isna().all()

def test_independent_formula_engine_financial_reconciliation():
    path=ROOT/'reports/formula_verification.json'
    if not path.exists():
        pytest.skip('Optional independent spreadsheet QA not run')
    verification=json.loads(path.read_text(encoding='utf-8'))
    assert verification['status']=='PASS'
    assert len(verification['financial_checks'])==3
    assert all(all(case['checks'].values()) for case in verification['financial_checks'])
